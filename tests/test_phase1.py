"""Phase 1: indicators, Momentum v1 rules, watchlist, results calendar, gates, scans."""
from __future__ import annotations

import json
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
import pytest

from agent import audit, settings
from agent.clock import IST
from agent.data import events, instruments, nse500, watchlist
from agent.db import connect, set_state
from agent.market_calendar import MarketCalendar
from agent.strategy import indicators, momentum_v1 as mv1, scanner

from .conftest import nse500_csv

STRAT = settings.load_strategy_config("momentum_v1")
P = STRAT.params


# ---------- indicators (reference values from TA-Lib 0.6, same seeded series) ----------

def _series(seed=11, n=260):
    rng = np.random.default_rng(seed)
    c = 100 * np.exp(np.cumsum(rng.normal(0, 0.02, n)))
    h = c * (1 + rng.uniform(0, 0.03, n))
    lo = c * (1 - rng.uniform(0, 0.03, n))
    return pd.DataFrame({"open": (h + lo) / 2, "high": h, "low": lo, "close": c, "volume": 1e5})


@pytest.mark.parametrize("bar, atr_ref, adx_ref, bb_ref", [
    (200, 3.980594, 21.992175, 119.461153),
    (259, 4.05076, 21.670236, 120.177407),
])
def test_indicators_match_talib(bar, atr_ref, adx_ref, bb_ref):
    df = _series()
    assert indicators.atr(df, 14).iloc[bar] == pytest.approx(atr_ref, abs=1e-4)
    assert indicators.adx(df, 14).iloc[bar] == pytest.approx(adx_ref, abs=0.02)
    assert indicators.bollinger_upper(df["close"], 20, 2.0).iloc[bar] == pytest.approx(bb_ref, abs=1e-6)


def test_warmup_rows_are_blank():
    out = indicators.compute(_series(n=150), P, warmup=100, tv_lookback=20)
    assert out.iloc[:99].isna().all().all()
    assert out.iloc[99:].notna().all().all()


# ---------- Momentum v1 levels and conditions ----------

def test_trade_levels_worked_example():
    bar = mv1.Bar("2026-09-22", 490, 505, 485, 500, 1e6, 495, 10, 25)
    lv = mv1.trade_levels(bar, [450.0], 0.05, P)
    assert (lv.entry_trigger, lv.entry_limit, lv.stop, lv.target) == (502.5, 505.05, 482.55, 560.0)
    assert lv.rr == pytest.approx((560 - 505.05) / (505.05 - 482.55), abs=1e-4)
    assert lv.rr_band == "B" and lv.rr_ok


def test_rr_bands_and_rejection():
    assert mv1.rr_band(2.5, P) == "A" and mv1.rr_band(2.49, P) == "B" and mv1.rr_band(1.5, P) == "C"
    assert mv1.rr_band(1.49, P) is None
    tight = mv1.trade_levels(mv1.Bar("d", 99, 101, 95, 100, 1, 99, 1, 25), [94.0], 0.05, P)
    assert not tight.rr_ok


def test_conditions():
    y = mv1.Bar("d0", 0, 0, 0, 100, 0, 101, 2.0, 22)
    t = mv1.Bar("d1", 0, 0, 0, 103, 0, 102, 2.1, 21)
    assert mv1.check_conditions(t, y, P).all_true
    assert not mv1.check_conditions(mv1.Bar("d1", 0, 0, 0, 102.9, 0, 102, 2.1, 21), y, P).change_ok
    assert not mv1.check_conditions(mv1.Bar("d1", 0, 0, 0, 103, 0, 104, 2.1, 21), y, P).bb_ok
    assert not mv1.check_conditions(mv1.Bar("d1", 0, 0, 0, 103, 0, 102, 2.0, 21), y, P).atr_rising
    assert not mv1.check_conditions(mv1.Bar("d1", 0, 0, 0, 103, 0, 102, 2.1, 20), y, P).adx_ok


def test_tick_rounding_is_conservative():
    assert mv1.round_up(100.01, 0.05) == 100.05 and mv1.round_down(100.04, 0.05) == 100.0
    assert mv1.round_up(100.05, 0.05) == 100.05


# ---------- watchlist ----------

def _universe(kite, symbols=("RELIANCE", "TCS", "INFY", "HDFCBANK")):
    instruments.sync_instruments()
    nse500.apply_constituents(nse500.parse_csv(nse500_csv([(s, "EQ", "IT") for s in symbols] + [("BEL", "BE", "Defence")])))


def test_watchlist_paste_formats(logged_in):
    _universe(logged_in)
    text = "NSE Code\tEPS\tProfit growth\nRELIANCE\t120.5\t25.3%\nTCS, 150, 18\nBEL,5.2,40\nZOMATO,1,50\nTCS,1,1\nbad line\n"
    res = watchlist.replace_watchlist(text).as_dict()
    assert res["symbols"] == ["RELIANCE", "TCS", "BEL-BE"]
    assert res["not_in_nse500"] == ["ZOMATO"] and res["duplicates"] == ["TCS"] and res["bad_lines"] == ["bad line"]
    assert watchlist.active()["RELIANCE"]["profit_growth"] == 25.3
    watchlist.replace_watchlist("INFY,10,20")
    assert list(watchlist.active()) == ["INFY"]


def test_watchlist_rejects_empty(logged_in):
    _universe(logged_in)
    with pytest.raises(watchlist.WatchlistFormatError):
        watchlist.replace_watchlist("hello world")


# ---------- results calendar ----------

class FakeNse:
    def __init__(self, rows=None, fail=False):
        self.rows, self.fail, self.calls = rows or [], fail, []

    def board_meetings(self, start, end):
        self.calls.append((start, end))
        if self.fail:
            raise RuntimeError("403 Forbidden")
        return [r for r in self.rows if start <= datetime.strptime(r["bm_date"], "%d-%b-%Y").date() <= end]


def _bm(sym, d, purpose="Financial Results"):
    return {"bm_symbol": sym, "bm_date": d.strftime("%d-%b-%Y"), "bm_purpose": purpose, "bm_desc": "Board Meeting Intimation"}


def test_events_store_only_results(fixed_clock, monkeypatch):
    fake = FakeNse([_bm("TCS", date(2026, 9, 25)), _bm("INFY", date(2026, 9, 26), "Fund Raising"),
                    _bm("RELIANCE", date(2026, 10, 1), "To consider Quarterly Unaudited Financial results and Dividend")])
    monkeypatch.setattr(events, "client_factory", lambda: fake)
    st = events.fetch_daily()
    assert st["ok"] and st["stored"] == 2
    assert events.status()["fetched_today"]


def test_events_failure_is_recorded_and_alerts(fixed_clock, monkeypatch):
    monkeypatch.setattr(events, "client_factory", lambda: FakeNse(fail=True))
    st = events.fetch_daily()
    assert not st["ok"] and not events.status()["fetched_today"]
    assert any(a["name"] == "results_calendar" for a in audit.active_alerts())


def test_results_check_windows(calendar):
    known = [d.isoformat() for d in pd.bdate_range("2026-09-01", "2026-09-25").date if calendar.is_trading_day(d)]
    d = date(2026, 9, 22)  # Tue; 5th session after = Tue 29 Sep
    assert scanner.results_check(d, [date(2026, 9, 29)], known, calendar, 5, 7)[0] is True
    assert scanner.results_check(d, [date(2026, 9, 22)], known, calendar, 5, 7)[0] is True   # today
    assert scanner.results_check(d, [date(2026, 9, 30)], known, calendar, 5, 7)[0] is False
    # beyond the cache the holiday calendar is used: from 30 Sep, Gandhi Jayanti (2 Oct) is skipped
    assert scanner.nth_session_after(date(2026, 9, 30), 2, known, calendar) == date(2026, 10, 5)
    # last year's same-quarter window: +/- 7 days
    assert scanner.results_check(d, [date(2025, 9, 29)], known, calendar, 5, 7)[0] is True
    assert scanner.results_check(d, [date(2025, 9, 30)], known, calendar, 5, 7)[0] is False


# ---------- end-to-end scan on crafted candles ----------

def _insert_candles(symbol, closes, highs, lows, vol=1e6, end=date(2026, 9, 22)):
    days = [d for d in pd.bdate_range(end=end, periods=len(closes) + 40).date
            if MarketCalendar(settings.get_app_config()).is_trading_day(d)][-len(closes):]
    with connect() as conn:
        conn.executemany("INSERT OR REPLACE INTO candles_daily (symbol, date, open, high, low, close, volume, fetched_at) "
                         "VALUES (?, ?, ?, ?, ?, ?, ?, 'x')",
                         [(symbol, d.isoformat(), c, h, lo, c, vol, ) for d, c, h, lo in zip(days, closes, highs, lows)])
    return days


def _breakout(n=160, jump=1.08):
    c = [100 * 1.003 ** t for t in range(n - 1)]
    c.append(c[-1] * jump)
    h = [x * 1.01 for x in c]
    lo = [x * 0.99 for x in c]
    h[-1], lo[-1] = c[-1] * 1.01, c[-2] * 0.995
    return c, h, lo


def _flat(n=160):
    c = [100.0 + (t % 3) * 0.1 for t in range(n)]
    return c, [x * 1.005 for x in c], [x * 0.995 for x in c]


@pytest.fixture
def market(logged_in, fixed_clock, monkeypatch):
    """Universe with a breakout in TCS on 22 Sep 2026, a flat INFY, Nifty above its 50-SMA."""
    fixed_clock.set(datetime(2026, 9, 22, 21, 0))
    _universe(logged_in)
    with connect() as conn:
        conn.execute("DELETE FROM candles_daily")
    _insert_candles("TCS", *_breakout())
    _insert_candles("INFY", *_flat())
    nifty = [20000 * 1.001 ** t for t in range(160)]
    _insert_candles("NIFTY 50", nifty, nifty, nifty, vol=0)
    indicators.compute_all(P)
    fake = FakeNse([])
    monkeypatch.setattr(events, "client_factory", lambda: fake)
    events.fetch_daily()
    return fake


def test_eod_scan_breakout_blocked_until_on_watchlist(market):
    s = scanner.scan_eod(STRAT)
    assert s["triggers"] == 1 and s["eligible"] == 0
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert sig["symbol"] == "TCS" and "Not on the watchlist" in sig["block_reasons"]
    ind = sig["indicators"]
    assert ind["change_pct"] >= 3 and ind["close"] >= ind["bb_upper"] and ind["adx14"] > 20
    assert ind["atr14"] > ind["atr14_prev"]

    close = ind["close"]
    watchlist.replace_watchlist(f"TCS,{close / 20:.4f},30")          # P/E 20 < growth 30
    s = scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert s["eligible"] == 1 and sig["eligible"] == 1, sig["block_reasons"]
    assert sig["gates"] == {"nse500": True, "watchlist": True, "valuation": True, "liquidity": True,
                            "results": True}                     # market filter off: not a gate

    watchlist.replace_watchlist(f"TCS,{close / 40:.4f},30")          # P/E 40 > growth 30
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert not sig["eligible"] and sig["block_reasons"][0].startswith("Valuation")


def test_eod_scan_results_meeting_blocks(market, monkeypatch):
    fake = FakeNse([_bm("TCS", date(2026, 9, 25))])
    monkeypatch.setattr(events, "client_factory", lambda: fake)
    events.fetch_daily()
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert sig["gates"]["results"] is False and "2026-09-25" in " ".join(sig["block_reasons"])


def test_eod_scan_fails_closed_without_results_calendar(market, monkeypatch):
    set_state(events.STATUS_KEY, {})                      # no good fetch on record
    monkeypatch.setattr(events, "client_factory", lambda: FakeNse(fail=True))
    events.fetch_daily()
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert not sig["eligible"] and "Results calendar not fetched since the 22 Sep close" in sig["block_reasons"]


def test_results_calendar_from_the_previous_evening_is_accepted(market, fixed_clock, monkeypatch):
    """The 22 Sep 21:00 fetch (after the close) covers 22 Sep signals scanned the next morning."""
    watchlist.replace_watchlist("TCS,10,1000")
    fixed_clock.set(datetime(2026, 9, 23, 9, 40))
    assert events.status()["fresh"] and not events.status()["fetched_today"]
    s = scanner.scan_eod(STRAT)
    assert s["date"] == "2026-09-22" and s["events_ok"] and s["eligible"] == 1
    # a failed refresh this morning keeps last evening's good data
    monkeypatch.setattr(events, "client_factory", lambda: FakeNse(fail=True))
    events.fetch_daily()
    assert events.fresh_for(date(2026, 9, 22)) and scanner.scan_eod(STRAT)["eligible"] == 1
    # after the 23 Sep close the same fetch is too old for 23 Sep signals
    fixed_clock.set(datetime(2026, 9, 23, 21, 0))
    assert not events.fresh_for(date(2026, 9, 23)) and not events.status()["fresh"]


def test_results_calendar_fetched_before_the_close_is_not_enough(market, fixed_clock):
    set_state(events.STATUS_KEY, {"daily": {"date": "2026-09-22", "ok": True, "at": "2026-09-22T11:00:00+05:30"}})
    fixed_clock.set(datetime(2026, 9, 23, 9, 40))
    assert not events.fresh_for(date(2026, 9, 22))


def test_rescan_keeps_signal_ids(market):
    """Trades and candidates point at signals by id: a re-scan must update rows in place."""
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    first = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    watchlist.replace_watchlist("INFY,10,1000")          # TCS drops off the watchlist
    scanner.scan_eod(STRAT)
    again = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert again["id"] == first["id"] and not again["eligible"]


def test_eod_scan_market_filter_blocks(market, market_filter_on):
    nifty = [20000 * 0.999 ** t for t in range(160)]
    _insert_candles("NIFTY 50", nifty, nifty, nifty, vol=0)
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert sig["gates"]["market"] is False and not sig["eligible"]


def test_eod_scan_ignores_nifty_when_filter_off(market):
    nifty = [20000 * 0.999 ** t for t in range(160)]
    _insert_candles("NIFTY 50", nifty, nifty, nifty, vol=0)          # Nifty well below its 50-SMA
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert "market" not in sig["gates"] and sig["eligible"], sig["block_reasons"]


def test_low_liquidity_blocks(market):
    c, h, lo = _breakout()
    _insert_candles("TCS", c, h, lo, vol=1000)   # ~₹2 lakh a day
    indicators.compute_all(P)
    watchlist.replace_watchlist("TCS,10,1000")
    scanner.scan_eod(STRAT)
    sig = scanner.signals_for("momentum_v1", day="2026-09-22")[0]
    assert sig["gates"]["liquidity"] is False


def test_history_scan_counts_by_month(market):
    summary = scanner.scan_history(STRAT)
    months = {m["month"]: m for m in summary["months"]}
    assert months["2026-09"]["triggers"] >= 1
    assert summary["totals"]["tradeable"] >= 1
    assert summary["results_history_from"] is None     # history not loaded: blackout not applied
    # EOD rows survive a history rerun
    scanner.scan_eod(STRAT)
    scanner.scan_history(STRAT)
    with connect() as conn:
        src = conn.execute("SELECT source FROM signals WHERE date='2026-09-22' AND symbol='TCS'").fetchone()[0]
    assert src == "eod"


def test_gate_table(market):
    watchlist.replace_watchlist("TCS,10,1000\nINFY,10,1")
    t = scanner.gate_table(only_watchlist=True)
    rows = {r["symbol"]: r for r in t["rows"]}
    assert rows["TCS"]["eligible"] and rows["TCS"]["gates"]["valuation"]
    assert not rows["INFY"]["gates"]["valuation"]          # P/E 10 > growth 1


# ---------- migration ----------

def test_migration_adds_signal_columns(tmp_path):
    import sqlite3
    from agent import db
    path = tmp_path / "old.db"
    with sqlite3.connect(path) as c:
        c.execute("CREATE TABLE signals (id INTEGER PRIMARY KEY AUTOINCREMENT, date TEXT NOT NULL, symbol TEXT NOT NULL, "
                  "strategy_version TEXT NOT NULL, entry_trigger REAL NOT NULL, entry_limit REAL NOT NULL, "
                  "stop REAL NOT NULL, target REAL NOT NULL, swing_low REAL, rr REAL NOT NULL, rr_band TEXT, "
                  "indicators TEXT, created_at TEXT NOT NULL, UNIQUE (date, symbol, strategy_version))")
    db.set_db_path(path)
    db.init_db(settings.get_app_config(), STRAT)
    with sqlite3.connect(path) as c:
        cols = {r[1] for r in c.execute("PRAGMA table_info(signals)")}
    assert {"eligible", "gates", "block_reasons", "source"} <= cols


# ---------- Screener results pasted into Excel ----------

def _screener_xlsx(rows, repeat_header_every=15):
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["S.No.", "Company", "CMP Rs.", "EPS 12M Rs.", "Profit growth %", "Mar Cap Rs.Cr. ",
              "Sales growth %", "Sales Var 3Yrs %"]
    ws.append(header)
    for i, (name, link_code, eps, growth) in enumerate(rows, 1):
        if i > 1 and (i - 1) % repeat_header_every == 0:
            ws.append(header)
        ws.append([i, name, 100.0, eps, growth, 5000, 10, 5])
        if link_code:
            ws.cell(ws.max_row, 2).hyperlink = f"https://www.screener.in/company/{link_code}/consolidated/"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_upload_reads_codes_from_links(logged_in):
    _universe(logged_in)
    rows = [("Reliance Industr", "RELIANCE", 120.5, 25.3), ("T C S", "TCS", 150, 18),
            ("Bharat Electron", "BEL", 5.2, 40), ("Tiny BSE Co", "532123", 1, 50),
            ("Not In Index", "ZOMATO", 1, 50), ("No Link Ltd", None, 1, 1)] + \
           [(f"Filler {i}", "INFY" if i == 0 else None, 10, 20) for i in range(12)]
    res = watchlist.replace_from_xlsx(_screener_xlsx(rows)).as_dict()
    assert res["symbols"] == ["RELIANCE", "TCS", "BEL-BE", "INFY"]
    assert res["bse_only"] == ["Tiny BSE Co"] and res["not_in_nse500"] == ["ZOMATO"]
    assert "No Link Ltd" in res["bad_lines"]               # a company name is not a code
    assert watchlist.active()["TCS"]["eps"] == 150


def test_excel_without_needed_columns_is_rejected(logged_in):
    import io
    import openpyxl
    _universe(logged_in)
    wb = openpyxl.Workbook()
    wb.active.append(["S.No.", "Company", "CMP Rs.", "P/E"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(watchlist.WatchlistFormatError, match="EPS 12M"):
        watchlist.replace_from_xlsx(buf.getvalue())


# ---------- gate 3: growth cap for weak-base spikes ----------

@pytest.mark.parametrize("pe, growth, ok", [
    (4.94, 2379.87, True),    # CPCL: cheap, passes whatever the growth
    (120.5, 135.8, False),    # expensive stock on a one-year spike: blocked by the 40 cap
    (36.6, 97.5, True),       # under the cap
    (40.0, 97.5, False),      # strictly below the cap
    (13.7, 14.3, True),       # growth below the cap: plain P/E < growth
    (37.1, 36.8, False),
    (-5.0, 50.0, False),      # loss-making
    (None, 50.0, False),
])
def test_valuation_growth_cap(pe, growth, ok):
    assert scanner.valuation_ok(pe, growth, 40.0)[0] is ok


def test_gate_table_applies_growth_cap(market):
    watchlist.replace_watchlist("TCS,1,300")
    tcs_close = scanner.gate_table(only_watchlist=True)["rows"][0]["close"]
    watchlist.replace_watchlist(f"TCS,{tcs_close / 60:.4f},300\nINFY,{100.1 / 20:.4f},300")  # P/E 60 and 20
    rows = {r["symbol"]: r for r in scanner.gate_table(only_watchlist=True)["rows"]}
    assert rows["TCS"]["gates"]["valuation"] is False and rows["TCS"]["growth_used"] == 40
    assert rows["INFY"]["gates"]["valuation"] is True
